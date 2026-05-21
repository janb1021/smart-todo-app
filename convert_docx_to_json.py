import docx
import json
import os
import re

def extract_text_from_docx(docx_path):
    doc = docx.Document(docx_path)
    content = []
    
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            style = paragraph.style.name if paragraph.style else "Normal"
            content.append({
                "type": "paragraph",
                "text": text,
                "style": style,
                "font_size": paragraph.style.font.size.pt if paragraph.style and paragraph.style.font.size else None,
                "bold": paragraph.runs[0].bold if paragraph.runs else None,
                "italic": paragraph.runs[0].italic if paragraph.runs else None
            })
    
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text.strip())
            if any(row_data):
                table_data.append(row_data)
        if table_data:
            content.append({
                "type": "table",
                "content": table_data
            })
    
    return content

def parse_exam_content(docx_path):
    doc = docx.Document(docx_path)
    exam_data = {
        "title": "",
        "sections": []
    }
    
    current_section = None
    current_question = None
    current_question_type = ""
    
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        
        if re.match(r'^2023年普通高等学校招生全国统一考试', text):
            exam_data["title"] = text
            continue
        
        match = re.match(r'^([一二三四五六七八九十]+)、(.+)', text)
        if match:
            if current_section:
                exam_data["sections"].append(current_section)
            current_section = {
                "section_number": match.group(1),
                "section_title": match.group(2),
                "questions": []
            }
            continue
        
        match = re.match(r'^(\d+)[．.、](.+)', text)
        if match and current_section:
            if current_question:
                current_section["questions"].append(current_question)
            question_num = match.group(1)
            question_text = match.group(2)
            
            if "选择题" in current_section.get("section_title", "") or "单项选择" in current_section.get("section_title", ""):
                current_question_type = "choice"
            elif "默写" in current_section.get("section_title", "") or "填空" in current_section.get("section_title", ""):
                current_question_type = "fill_in_blank"
            elif "阅读" in current_section.get("section_title", "") or "现代文" in current_section.get("section_title", "") or "古诗文" in current_section.get("section_title", ""):
                current_question_type = "reading"
            elif "作文" in current_section.get("section_title", ""):
                current_question_type = "composition"
            else:
                current_question_type = "other"
            
            current_question = {
                "question_number": question_num,
                "question_type": current_question_type,
                "content": question_text,
                "options": [],
                "answer": "",
                "analysis": ""
            }
            continue
        
        if current_question and current_question_type == "choice":
            option_match = re.match(r'^([A-D])[．.、](.+)', text)
            if option_match:
                current_question["options"].append({
                    "option": option_match.group(1),
                    "content": option_match.group(2)
                })
                continue
        
        if current_question and ("答案" in text or "【答案】" in text or "参考答案" in text):
            current_question["answer"] = text
            continue
        
        if current_question and ("解析" in text or "【解析】" in text):
            current_question["analysis"] = text
            continue
        
        if current_question:
            current_question["content"] += "\n" + text
    
    if current_question:
        current_section["questions"].append(current_question)
    if current_section:
        exam_data["sections"].append(current_section)
    
    return exam_data

def main():
    docx_path = r"C:\Users\janb\Downloads\2023年普通高等学校招生全国统一考试语文+答案.docx"
    
    if not os.path.exists(docx_path):
        print(f"错误：文件不存在 - {docx_path}")
        return
    
    print(f"正在读取Word文件: {docx_path}")
    
    exam_data = parse_exam_content(docx_path)
    
    json_path = os.path.join(os.path.dirname(docx_path), "2023年普通高等学校招生全国统一考试语文+答案.json")
    
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(exam_data, f, ensure_ascii=False, indent=2)
    
    print(f"JSON文件已保存: {json_path}")
    
    print("\n转换结果摘要:")
    print(f"标题: {exam_data['title']}")
    print(f"章节数量: {len(exam_data['sections'])}")
    total_questions = sum(len(section['questions']) for section in exam_data['sections'])
    print(f"题目总数: {total_questions}")
    
    for section in exam_data['sections']:
        print(f"  - {section['section_number']}、{section['section_title']}: {len(section['questions'])}题")

if __name__ == "__main__":
    main()